"""주간 메뉴표 엑셀 파서."""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.worksheet import Worksheet

MenuItem = tuple[str, str, str]
ParsedMenus = dict[date, dict[str, list[MenuItem]]]

MEAL_LABELS = {
    "중식": "lunch",
    "점심": "lunch",
    "석식": "dinner",
    "저녁": "dinner",
}
SKIP_LABELS = {
    "구분",
    "주간메뉴표",
    "kcal",
    "칼로리",
    "비고",
    "원산지",
    "열량",
}
SKIP_ITEMS = {"", "-", "/", "미운영", "(미운영)"}


def _compact(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).strip()


def _as_date(value: Any) -> date | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if value < 20000 or value > 80000:
            return None
        try:
            parsed = from_excel(value)
        except Exception:
            return None
        if isinstance(parsed, datetime):
            return parsed.date()
        if isinstance(parsed, date):
            return parsed
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace("년", ".").replace("월", ".").replace("일", "")
    text = re.sub(r"[./]\s*$", "", text)
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%y.%m.%d", "%y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _split_menu_text(raw: Any) -> tuple[str, str] | None:
    if raw is None:
        return None
    text = str(raw).strip()
    compact = _compact(text)
    if compact in SKIP_ITEMS:
        return None
    origins = re.findall(r"\[([^\]]+)\]", text)
    name = re.sub(r"\s*\[[^\]]+\]", "", text)
    name = re.sub(r"[ \t]+", " ", name).replace("\n", " ").strip()
    if not name or _compact(name) in SKIP_ITEMS:
        return None
    return name, ", ".join(origins)


def _merge_lookup(ws: Worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    lookup: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        origin = (rng.min_row, rng.min_col)
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                lookup[(row, col)] = origin
    return lookup


def _cell(ws: Worksheet, row: int, col: int, lookup: dict[tuple[int, int], tuple[int, int]]):
    origin = lookup.get((row, col), (row, col))
    return ws.cell(origin[0], origin[1]).value


def _date_columns(ws: Worksheet, row: int, lookup) -> list[tuple[int, date]]:
    found: list[tuple[int, date]] = []
    for col in range(1, (ws.max_column or 1) + 1):
        parsed = _as_date(_cell(ws, row, col, lookup))
        if parsed:
            found.append((col, parsed))
    return found


def _load_workbook(source: str | Path | BytesIO | BinaryIO):
    if isinstance(source, (str, Path)):
        return load_workbook(source, data_only=True)
    data = source.read()
    if hasattr(source, "seek"):
        try:
            source.seek(0)
        except Exception:
            pass
    return load_workbook(BytesIO(data), data_only=True)


def parse_menu_sheet(ws: Worksheet) -> ParsedMenus:
    lookup = _merge_lookup(ws)
    max_row = ws.max_row or 1
    header_rows: list[tuple[int, list[tuple[int, date]]]] = []
    for row in range(1, max_row + 1):
        cols = _date_columns(ws, row, lookup)
        if len(cols) >= 1 and any(_as_date(_cell(ws, row, col, lookup)) for col, _ in cols):
            # 날짜가 2개 이상이면 주간 헤더로 본다. 1개만 있어도 한 줄 식단일 수 있다.
            if len(cols) >= 1:
                header_rows.append((row, cols))

    # 같은 주의 병합된 날짜 행(3행·4행)은 한 번만 쓴다.
    unique_headers: list[tuple[int, list[tuple[int, date]]]] = []
    seen_dates: set[tuple[date, ...]] = set()
    for row, cols in header_rows:
        key = tuple(day for _, day in cols)
        if key in seen_dates:
            continue
        seen_dates.add(key)
        unique_headers.append((row, cols))

    result: ParsedMenus = {}
    for idx, (header_row, date_cols) in enumerate(unique_headers):
        end_row = unique_headers[idx + 1][0] if idx + 1 < len(unique_headers) else max_row + 1
        meal_type = ""
        section = "중식"
        for row in range(header_row + 1, end_row):
            label = _compact(_cell(ws, row, 1, lookup))
            if label in SKIP_LABELS or label.lower() in SKIP_LABELS:
                continue
            if label in MEAL_LABELS:
                meal_type = MEAL_LABELS[label]
                section = "중식" if meal_type == "lunch" else "석식"
            elif label.upper() == "PLUSBAR":
                section = "PLUS BAR"
            elif not meal_type:
                continue

            if not meal_type:
                continue
            for col, day in date_cols:
                bucket = result.setdefault(day, {})
                items = bucket.setdefault(meal_type, [])
                parsed = _split_menu_text(_cell(ws, row, col, lookup))
                if not parsed:
                    continue
                name, description = parsed
                item = (section, name, description)
                if item not in items:
                    items.append(item)
    return result


def parse_menu_workbook(source: str | Path | BytesIO | BinaryIO) -> ParsedMenus:
    wb = _load_workbook(source)
    merged: ParsedMenus = {}
    try:
        for name in wb.sheetnames:
            parsed = parse_menu_sheet(wb[name])
            for day, meals in parsed.items():
                slot = merged.setdefault(day, {})
                for meal_type, items in meals.items():
                    existing = slot.setdefault(meal_type, [])
                    for item in items:
                        if item not in existing:
                            existing.append(item)
    finally:
        wb.close()
    if not any(items for meals in merged.values() for items in meals.values()):
        raise ValueError("식단표에서 메뉴를 찾지 못했습니다. 주간 메뉴표 양식(날짜 + 중식/석식)인지 확인해 주세요.")
    return merged


def import_menus_from_workbook(conn, source: str | Path | BytesIO | BinaryIO) -> dict[str, Any]:
    from db import replace_menus

    parsed = parse_menu_workbook(source)
    meal_slots = 0
    item_count = 0
    for day, meals in parsed.items():
        for meal_type, items in meals.items():
            replace_menus(conn, day, meal_type, items)
            if items:
                meal_slots += 1
                item_count += len(items)
    days = sorted(parsed)
    return {
        "days": days,
        "meal_slots": meal_slots,
        "item_count": item_count,
        "start": days[0] if days else None,
        "end": days[-1] if days else None,
    }
